import os
import csv
from docx import Document

def read_keywords(reference_file):
    keywords = []
    file_extension = os.path.splitext(reference_file)[1].lower()

    try:
        if file_extension == '.csv':
            with open(reference_file, 'r', encoding='utf-8') as file:
                reader = csv.reader(file)
                keywords = [row[0].strip().lower() for row in reader if row]
        elif file_extension == '.txt':
            with open(reference_file, 'r', encoding='utf-8') as file:
                keywords = [line.strip().lower() for line in file if line.strip()]
        else:
            raise ValueError(f"Unsupported file format: {file_extension}")
    except Exception as e:
        raise Exception(f"Error reading reference file: {str(e)}")

    return keywords

def tag_document(doc_path, keywords):
    doc = Document(doc_path)
    text = " ".join([para.text.lower() for para in doc.paragraphs])

    tags = [keyword for keyword in keywords if keyword in text]

    if tags:
        doc.add_paragraph()
        tag_paragraph = doc.add_paragraph("Tags: ")
        tag_paragraph.add_run(', '.join(f'"{tag}"' for tag in tags))

        doc.save(doc_path)

    return tags