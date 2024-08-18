import logging
from collections import Counter
from docx import Document
from openpyxl import Workbook
import os

logger = logging.getLogger(__name__)

def create_word_count_summary(doc_paths, output_folder, min_count=20, max_count=100):
    try:
        word_count = Counter()

        for doc_path in doc_paths:
            try:
                doc = Document(doc_path)
                for para in doc.paragraphs:
                    words = para.text.lower().split()
                    word_count.update(words)
            except Exception as e:
                logger.error(f"Error processing document {doc_path}: {str(e)}")

        filtered_words = {word: count for word, count in word_count.items() if min_count <= count <= max_count}

        summary_filename = f"word_count_summary_{min_count}_to_{max_count}.xlsx"
        summary_path = os.path.join(output_folder, summary_filename)

        wb = Workbook()
        ws = wb.active
        ws.title = "Word Count Summary"
        ws.append(["Word", "Count"])

        if filtered_words:
            for word, count in sorted(filtered_words.items(), key=lambda x: x[1], reverse=True):
                ws.append([word, count])
            message = f"Word count summary created with {len(filtered_words)} words"
            logger.info(message)
        else:
            min_count_found = min(word_count.values()) if word_count else 0
            max_count_found = max(word_count.values()) if word_count else 0
            ws.append(["No words found in the specified range"])
            ws.append([f"Word count range in documents: {min_count_found} - {max_count_found}"])
            message = f"No words found with count between {min_count} and {max_count}. Word count range in documents: {min_count_found} - {max_count_found}"
            logger.warning(message)

        wb.save(summary_path)
        logger.info(f"Summary file saved: {summary_path}")

        return summary_path, message
    except Exception as e:
        error_message = f"Error creating word count summary: {str(e)}"
        logger.exception(error_message)
        return None, error_message